# 原告信息表的读写：解析律师维护的 Excel、规范化、校验、比对变更、存档与回退。
#
# 表是全公司共用的底表（不是单个案件的材料），由管理员在「原告信息表」页面上传更新，
# 存在服务器 backend/data/ 下作为唯一数据来源（不进 git）。每次更新把旧版本存进历史目录，
# 传错了可以回退。查询（按原告名称取信用代码 / 负责人 / 住址）见 branch_registry.py。
#
# 存储格式（branch_info.json）：
#   {"updated_at": "2026-09-29T10:00:00+08:00", "source_file": "原告信息表.xls",
#    "entries": [{"name", "credit_code", "address", "representative", "phone"}]}
# 兼容早期手写格式 {"headquarters": {...}, "branches": [...]}（读取时转换）。
from __future__ import annotations

import io
import json
import os
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone

from app.services.validators import validate_credit_code

# 表中各列的字段名 → 中文列名（导出 Excel 用同一套表头，律师改完可直接传回）
COLUMNS: list[tuple[str, str]] = [
    ("name", "原告名称"),
    ("credit_code", "统一社会信用代码"),
    ("address", "住所地"),
    ("representative", "法定代表人\\负责人"),
    ("phone", "联系方式"),
]
COLUMN_LABELS = dict(COLUMNS)

# 表头识别：按关键词匹配，律师表头写法略有出入（「住址」「负责人」等）也能认出
_HEADER_KEYWORDS: dict[str, tuple[str, ...]] = {
    "name": ("原告名称", "公司名称", "单位名称", "名称"),
    "credit_code": ("信用代码",),
    "address": ("住所", "住址", "地址"),
    "representative": ("法定代表人", "负责人"),
    "phone": ("联系方式", "电话"),
}

_CN_TZ = timezone(timedelta(hours=8))
_HALF_TO_FULL = str.maketrans({"(": "（", ")": "）"})


@dataclass
class BranchTable:
    """当前生效的原告信息表。"""

    entries: list[dict[str, str]]
    updated_at: str | None = None
    source_file: str | None = None


@dataclass
class CheckResult:
    """上传表的校验结论：errors 阻止保存，warnings 仅提示（如信用代码校验位不符）。"""

    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def normalize_name(name: str) -> str:
    """公司名规范化：去空白、半角括号转全角（系统拼出的原告名称一律用全角括号）。"""
    return re.sub(r"\s+", "", name or "").translate(_HALF_TO_FULL)


def _clean(value: object) -> str:
    text = "" if value is None else str(value).strip()
    # xlrd 把纯数字单元格读成浮点（如电话 83679974.0），去掉多余的 .0
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def normalize_entry(raw: dict[str, object]) -> dict[str, str]:
    entry = {key: _clean(raw.get(key)) for key, _ in COLUMNS}
    entry["name"] = normalize_name(entry["name"])
    entry["credit_code"] = entry["credit_code"].replace(" ", "").upper()
    return entry


# ── 解析 Excel ──────────────────────────────────────────────────────────────


def _read_rows(filename: str, content: bytes) -> list[list[object]]:
    """读出第一个非空工作表的所有行。支持 .xls（WPS/老版 Excel）与 .xlsx。"""
    lower = filename.lower()
    try:
        if lower.endswith(".xls"):
            import xlrd

            book = xlrd.open_workbook(file_contents=content)
            for sheet in book.sheets():
                if sheet.nrows:
                    return [sheet.row_values(r) for r in range(sheet.nrows)]
            return []
        if lower.endswith(".xlsx"):
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            for sheet in book.worksheets:
                rows = [list(r) for r in sheet.iter_rows(values_only=True)]
                if any(any(c not in (None, "") for c in r) for r in rows):
                    return rows
            return []
    except Exception as e:  # 损坏文件、加密文件等：给出明确中文错误，不静默失败
        raise ValueError(f"无法读取表格文件：{e}") from e
    raise ValueError("仅支持 .xls 或 .xlsx 格式的表格")


def _match_header(row: list[object]) -> dict[str, int] | None:
    """在一行里找各列位置；认出「名称」和「信用代码」两列才算表头。"""
    cols: dict[str, int] = {}
    for idx, cell in enumerate(row):
        text = re.sub(r"\s+", "", _clean(cell))
        for key, words in _HEADER_KEYWORDS.items():
            if key not in cols and any(w in text for w in words):
                cols[key] = idx
                break
    return cols if {"name", "credit_code"} <= cols.keys() else None


def parse_table(filename: str, content: bytes) -> list[dict[str, str]]:
    """Excel → 规范化后的条目列表（跳过空行）。找不到表头时抛 ValueError。"""
    rows = _read_rows(filename, content)
    for i, row in enumerate(rows[:10]):  # 表头一般在前几行（上面可能有标题行）
        cols = _match_header(row)
        if cols is None:
            continue
        entries = []
        for data in rows[i + 1 :]:
            raw = {k: data[c] if c < len(data) else "" for k, c in cols.items()}
            entry = normalize_entry(raw)
            if any(entry.values()):
                entries.append(entry)
        return entries
    raise ValueError("未找到表头：表格需包含「原告名称」和「统一社会信用代码」两列")


# ── 校验与比对 ──────────────────────────────────────────────────────────────


def check_entries(entries: list[dict[str, str]]) -> CheckResult:
    """缺名称 / 重名 / 空表 → 错误；信用代码不合法、缺负责人或住址 → 提醒。"""
    result = CheckResult()
    if not entries:
        result.errors.append("表格中没有任何数据行")
        return result
    seen: set[str] = set()
    for n, e in enumerate(entries, start=1):
        name = e.get("name", "")
        if not name:
            result.errors.append(f"第 {n} 行缺少原告名称")
            continue
        if name in seen:
            result.errors.append(f"「{name}」重复出现")
        seen.add(name)
        check = validate_credit_code(e.get("credit_code"))
        if not check.passed:
            result.warnings.append(f"「{name}」{check.message}")
        for key in ("address", "representative"):
            if not e.get(key):
                result.warnings.append(f"「{name}」缺少{COLUMN_LABELS[key]}")
    return result


def diff_entries(
    old: list[dict[str, str]], new: list[dict[str, str]]
) -> list[dict[str, object]]:
    """按名称比对新旧两版：新增 / 删除 / 修改（列出每个改动的列）。"""
    old_map = {e["name"]: e for e in old}
    new_map = {e["name"]: e for e in new}
    diff: list[dict[str, object]] = []
    for name, entry in new_map.items():
        before = old_map.get(name)
        if before is None:
            diff.append({"name": name, "kind": "added", "changes": []})
            continue
        changes = [
            {
                "field": key,
                "label": label,
                "old": before.get(key, ""),
                "new": entry[key],
            }
            for key, label in COLUMNS
            if key != "name" and before.get(key, "") != entry[key]
        ]
        if changes:
            diff.append({"name": name, "kind": "changed", "changes": changes})
    for name in old_map.keys() - new_map.keys():
        diff.append({"name": name, "kind": "removed", "changes": []})
    return diff


# ── 读写与历史版本 ────────────────────────────────────────────────────────────


def _from_legacy(data: dict) -> list[dict[str, str]]:
    """早期手写格式 {headquarters, branches} → 条目列表。"""
    items = []
    if isinstance(data.get("headquarters"), dict):
        items.append(data["headquarters"])
    items.extend(b for b in data.get("branches") or [] if isinstance(b, dict))
    return [
        normalize_entry(
            {**b, "representative": b.get("person_in_charge") or b.get("legal_rep")}
        )
        for b in items
    ]


def load_table(path: str) -> BranchTable | None:
    """读当前表；文件不存在或损坏返回 None（调用方把原告几项留缺失，不猜）。"""
    if not os.path.exists(path):
        return None
    try:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
    except (OSError, json.JSONDecodeError):
        return None
    if not isinstance(data, dict):
        return None
    if "entries" in data:
        entries = [normalize_entry(e) for e in data["entries"] if isinstance(e, dict)]
    else:
        entries = _from_legacy(data)
    return BranchTable(entries, data.get("updated_at"), data.get("source_file"))


def _write_json(path: str, table: BranchTable) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    tmp = f"{path}.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(
            {
                "updated_at": table.updated_at,
                "source_file": table.source_file,
                "entries": table.entries,
            },
            f,
            ensure_ascii=False,
            indent=2,
        )
    os.replace(tmp, path)  # 原子替换，避免写到一半时被查询读到残缺文件


def _history_id(table: BranchTable) -> str:
    stamp = table.updated_at or datetime.now(_CN_TZ).isoformat()
    return re.sub(r"[^0-9]", "", stamp)[:14] or "legacy"


def save_table(
    entries: list[dict[str, str]],
    source_file: str | None,
    path: str,
    history_dir: str,
    now: datetime | None = None,
) -> BranchTable:
    """保存新版本为当前表；保存前把旧版本存入历史目录。"""
    current = load_table(path)
    if current is not None:
        _write_json(os.path.join(history_dir, f"{_history_id(current)}.json"), current)
    stamp = (now or datetime.now(_CN_TZ)).isoformat(timespec="seconds")
    table = BranchTable([normalize_entry(e) for e in entries], stamp, source_file)
    _write_json(path, table)
    return table


def list_history(history_dir: str) -> list[dict[str, object]]:
    """历史版本列表（新 → 旧）。"""
    if not os.path.isdir(history_dir):
        return []
    items = []
    for fname in sorted(os.listdir(history_dir), reverse=True):
        if not fname.endswith(".json"):
            continue
        table = load_table(os.path.join(history_dir, fname))
        if table is not None:
            items.append(
                {
                    "id": fname[:-5],
                    "updated_at": table.updated_at,
                    "source_file": table.source_file,
                    "count": len(table.entries),
                }
            )
    return items


def restore_table(
    version_id: str, path: str, history_dir: str, now: datetime | None = None
) -> BranchTable:
    """把某个历史版本恢复为当前表（恢复本身也是一次更新，当前版本同样先存档）。"""
    if not re.fullmatch(r"[0-9A-Za-z_-]+", version_id):
        raise ValueError("版本号不合法")
    old = load_table(os.path.join(history_dir, f"{version_id}.json"))
    if old is None:
        raise ValueError("找不到该历史版本")
    source = (
        f"恢复自 {old.updated_at or version_id} 版（{old.source_file or '未知文件'}）"
    )
    return save_table(old.entries, source, path, history_dir, now)


def age_days(table: BranchTable, now: datetime | None = None) -> int | None:
    """距上次更新的天数；没有更新时间（早期手写格式）返回 None。"""
    if not table.updated_at:
        return None
    try:
        updated = datetime.fromisoformat(table.updated_at)
    except ValueError:
        return None
    if updated.tzinfo is None:
        updated = updated.replace(tzinfo=_CN_TZ)
    return max(0, ((now or datetime.now(_CN_TZ)) - updated).days)


def to_xlsx(entries: list[dict[str, str]]) -> bytes:
    """导出为 .xlsx（表头与上传格式一致，律师改完可直接传回）。"""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "原告信息表"
    sheet.append([label for _, label in COLUMNS])
    for e in entries:
        sheet.append([e.get(key, "") for key, _ in COLUMNS])
    for col, width in zip("ABCDE", (40, 24, 60, 18, 16)):
        sheet.column_dimensions[col].width = width
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()
